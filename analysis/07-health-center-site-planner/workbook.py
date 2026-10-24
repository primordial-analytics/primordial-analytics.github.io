"""The downloadable screen: type a ZIP and a capacity, get the draw, its source split and payer mix. Live formulas, no macros.

The ZIP table sheet holds, for every ZCTA, the share of a typical single site's patients at the centroid who have no health center
today (10th, 50th and 90th percentile over the 24 bootstrap parameter sets the page uses) and the payer mix. Neither depends on capacity. The Screen sheet
multiplies them by the capacity the user types: capacity is an input, the model does not forecast how many patients a site draws.
"""
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, GOLD = "191345", "B3A369"
HEAD = PatternFill("solid", fgColor=GOLD)
INPUT = PatternFill("solid", fgColor="FFF2B3")
CALC = PatternFill("solid", fgColor="EFEFEF")
BOLD = Font(bold=True, color=NAVY)
TITLE = Font(bold=True, size=16, color=NAVY)
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
ROWS = 25
FIRST = 6

TABLE_COLS = ["zcta", "state", "county", "density_class", "low_income_residents", "patients_all_centers", "penetration_p50", "not_patients_p10",
              "not_patients_p50", "not_patients_p90", "hrsa_published_penetration", "small_zip", "nearest_site_miles", "site_share_new_p10", "site_share_new_p50",
              "site_share_new_p90", "site_capacity_binds", "site_payer_uninsured", "site_payer_medicaid", "site_payer_medicare", "site_payer_private", "campus_zip"]


def add_cached_values(path, skip=("ZIP table",)):
    """openpyxl writes formulas without results, so a viewer that does not calculate shows blank cells. Recalculate a copy in
    LibreOffice (headless) and write its results into the original file, which keeps our formatting and validation."""
    soffice = shutil.which("soffice")
    if not soffice:
        print("soffice not found: workbook saved without cached values")
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)], check=True, capture_output=True)
        calc = load_workbook(Path(tmp) / path.name, data_only=True)
    names = load_workbook(path, read_only=True).sheetnames
    cell = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: [a-z]+="[^"]*")*)><f>(.*?)</f><v ?/></c>', re.S)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    for i, name in enumerate(names, start=1):
        if name in skip:
            continue
        ws = calc[name]

        def fill(m):
            v = ws[m.group(1)].value
            if v is None:
                return m.group(0)
            if isinstance(v, str):
                return f'<c r="{m.group(1)}"{m.group(2)} t="str"><f>{m.group(3)}</f><v>{escape(v)}</v></c>'
            return f'<c r="{m.group(1)}"{m.group(2)}><f>{m.group(3)}</f><v>{v!r}</v></c>'

        key = f"xl/worksheets/sheet{i}.xml"
        members[key] = cell.sub(fill, members[key].decode()).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in members.items():
            z.writestr(n, data)
    return True


def put(ws, ref, value, fmt=None, fill=None, bold=False):
    cell = ws[ref]
    cell.value = value
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
        cell.border = BOX
    if bold:
        cell.font = BOLD
    return cell


def table_sheet(wb, table):
    ws = wb.create_sheet("ZIP table")
    cols = [c for c in TABLE_COLS]
    ws.append(cols)
    t = table.reset_index(names="zcta")[cols]
    for row in t.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and v != v) else v for v in row])
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).font = BOLD
        ws.cell(1, c).fill = HEAD
        ws.column_dimensions[get_column_letter(c)].width = 14
    for r in range(2, len(t) + 2):
        ws.cell(r, 1).number_format = "@"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(t) + 1}"
    return len(t) + 1, {c: get_column_letter(i + 1) for i, c in enumerate(cols)}


def make(table, example, facts, out_dir, slug, cap_default, cap_p25, cap_p75, k_screen, for_check=None):
    """Build the workbook. for_check: optional list of (zip, capacity) rows to fill instead of the worked example (used to test it)."""
    out_dir = Path(out_dir)
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    sc = wb.create_sheet("Screen")
    notes = wb.create_sheet("Assumptions and sources")
    last, col = table_sheet(wb, table)

    def look(name, row_ref):
        return f"INDEX('ZIP table'!${col[name]}$2:${col[name]}${last},{row_ref})"

    # ---------------- Screen
    ws = sc
    ws["A1"] = "Site screen: who a typical single site's patients would be in a ZIP code"
    ws["A1"].font = TITLE
    ws["A2"] = "Type a ZIP code (a ZCTA) and the patients a year the site could see: that number is your assumption. Yellow cells are yours. The first row is a worked example."
    ws["A3"] = "A scenario for a typical center on public data, not a volume forecast. Ranges are 10th to 90th percentile over 24 refits of the model (the same 24 the page uses), and cover model and survey error only."
    heads = [("ZIP code", 9), ("Capacity (patients a year)", 12), ("State", 7), ("County", 20), ("Low-income residents", 12), ("Patients of all centers today", 12),
             ("Low-income residents who are not patients (likely)", 14), ("Same, low end", 11), ("Same, high end", 11), ("Share who have no center today, low", 11),
             ("Share who have no center today, likely", 11), ("Share who have no center today, high", 11), ("New patients (no center today), low", 11),
             ("New patients (no center today), likely", 11), ("New patients (no center today), high", 11), ("Patients taken from neighbours (likely)", 11),
             ("Uninsured", 9), ("Medicaid, CHIP, other public", 11), ("Medicare", 9), ("Private", 9), ("Row in ZIP table", 8), ("Note", 52)]
    for i, (h, w) in enumerate(heads, start=1):
        c = ws.cell(FIRST - 1, i, h)
        c.font = BOLD
        c.fill = HEAD
        c.alignment = WRAP
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[FIRST - 1].height = 58
    ex_rows = for_check or [(example["zcta"], example["capacity"])]
    n_rows = max(ROWS, len(ex_rows))
    for r in range(FIRST, FIRST + n_rows):
        i = r - FIRST
        a, b = f"$A{r}", f"$B{r}"
        if i < len(ex_rows):
            put(ws, f"A{r}", ex_rows[i][0], "@", INPUT)
            put(ws, f"B{r}", ex_rows[i][1], "#,##0", INPUT)
        else:
            put(ws, f"A{r}", None, "@", INPUT)
            put(ws, f"B{r}", None, "#,##0", INPUT)
        m = f"$U{r}"
        put(ws, f"U{r}", f'=IF({a}="","",IFERROR(MATCH(TEXT({a},"00000"),\'ZIP table\'!$A$2:$A${last},0),"not found"))', fill=CALC)
        ok = f'ISNUMBER({m})'
        cells = {
            "C": (look("state", m), None), "D": (look("county", m), None), "E": (look("low_income_residents", m), "#,##0"),
            "F": (look("patients_all_centers", m), "#,##0"), "G": (look("not_patients_p50", m), "#,##0"), "H": (look("not_patients_p10", m), "#,##0"),
            "I": (look("not_patients_p90", m), "#,##0"), "J": (look("site_share_new_p10", m), "0%"), "K": (look("site_share_new_p50", m), "0%"),
            "L": (look("site_share_new_p90", m), "0%"), "M": (f"{b}*J{r}", "#,##0"), "N": (f"{b}*K{r}", "#,##0"), "O": (f"{b}*L{r}", "#,##0"),
            "P": (f"{b}-N{r}", "#,##0"),
            "Q": (f"{b}*{look('site_payer_uninsured', m)}", "#,##0"), "R": (f"{b}*{look('site_payer_medicaid', m)}", "#,##0"),
            "S": (f"{b}*{look('site_payer_medicare', m)}", "#,##0"), "T": (f"{b}*{look('site_payer_private', m)}", "#,##0"),
        }
        # shares and ZIP facts always show; patient counts only where a typical site here would fill the capacity (the page caps counts at the model's draw)
        valid = f'IF({ok},IF(ISNUMBER({b}),AND({b}<={k_screen},{look("site_capacity_binds", m)}=1),FALSE),FALSE)'
        for c, (f, fmt) in cells.items():
            put(ws, f"{c}{r}", f'=IF({valid if c in "MNOPQRST" else ok},{f},"")', fmt, CALC)
        put(ws, f"V{r}", (f'=IF({a}="","",IF(NOT({ok}),"ZIP not in the table: use a 5-digit ZIP code tabulation area.",'
                          f'IF({look("small_zip", m)}=1,"Fewer than 1,000 low-income residents: ACS margins are wide here. ","")&'
                          f'IF({look("campus_zip", m)}=1,"A campus ZIP: the Census counts college students as low-income. ","")&'
                          f'IF({b}>{k_screen},"Above {k_screen:,} patients a year the screen cannot say whether a typical site here would fill this capacity, so patient counts are blank. The shares apply.",'
                          f'IF({look("site_capacity_binds", m)}=1,"A typical site here would draw more than this capacity, so capacity sets the patients and the counts are shown.",'
                          f'"A typical site here may not fill this capacity, so patient counts are blank. The shares apply."))))'), fill=CALC)
        ws[f"V{r}"].alignment = WRAP
    dv = DataValidation(type="whole", operator="between", formula1="100", formula2="15000", allow_blank=True, showErrorMessage=True, errorStyle="stop",
                        errorTitle="Capacity", error="Enter a whole number of patients a year from 100 to 15,000.")
    ws.add_data_validation(dv)
    dv.add(f"B{FIRST}:B{FIRST + n_rows - 1}")
    ws.freeze_panes = f"C{FIRST}"
    ws.cell(FIRST + n_rows + 1, 1, f"Typical site capacity for reference: median {cap_default:,} patients a year per site (25th percentile {cap_p25:,}, 75th percentile {cap_p75:,}).")

    # ---------------- Assumptions and sources
    ws = notes
    ws.column_dimensions["A"].width = 120
    ws["A1"] = "What this is, and what it is not"
    ws["A1"].font = TITLE
    lines = [
        "An example on public data, not client work. No patient information is used. The numbers are a scenario for a typical center, not a forecast of patient volume.",
        "",
        "How the shares are made. HRSA reports how many patients each health center has in each ZIP code (2025). We fit a model of where a center's patients come from: closer sites draw more, "
        "a center takes a smaller share where others already compete, and some people use no health center at all. The ZIP table holds, for a typical single site placed at the ZIP's centroid, the share of "
        "its patients who have no health center today and the payer mix. Neither depends on capacity.",
        f"Capacity is your input. The model does not say how many patients a site sees: for a typical single site it gives more than {k_screen:,} a year in most ZIP codes, more than real single sites see, so volume depends on staffing, hours and rooms, which are not in the data. "
        "The Screen sheet multiplies the shares by the capacity you type, from 100 to 15,000 patients a year. "
        f"It shows patient counts only where the model's draw for a typical site is above {k_screen:,} a year (the 'site_capacity_binds' column of the ZIP table) and your capacity is {k_screen:,} or less: there capacity, not the model, sets the number of patients, as on the page. "
        "Elsewhere a site might not fill your capacity, so the counts are blank and only the shares show.",
        "Who they are. The share labelled 'no center today' comes from people who use no health center now; the rest is taken from neighbouring centers.",
        "Payer mix. Patients taken from neighbours keep the observed payer mix of their ZIP. New patients are uninsured at the ACS share of low-income residents of all ages in that ZIP; the insured remainder follows the ZIP's observed Medicaid, Medicare and private pattern.",
        "Ranges. The low and high ends of the shares are the 10th and 90th percentile over 24 refits of the model (a bootstrap over centers), the same 24 the page uses, so the shares agree with the page for the same ZIP, and so do the counts where they are shown. "
        "The ranges of people not reached come from the ACS margins of error and the censored cells. Patient-count noise in a single year is not included.",
        "",
        "Limits. Straight-line miles, not drive time. A typical center: staffing, hours, language and referrals are not in the data. Patients' income is scaled from each center's overall mix. "
        "Look-alike centers are competitors, but their patients are not in the UDS file. A person served by two centers is counted twice. Seasonal, mobile and school sites are not in the supply. "
        "Data years differ: UDS 2025, ACS 2020-2024, a site list from October 2026. People who are not health center patients may get care elsewhere.",
        "",
        "Sources: HRSA Uniform Data System 2025 (patients by ZIP and payer, Table 4), HRSA health center site list (October 2026), Census ACS 2020-2024 5-year tables C17002 and C27016, "
        "Census 2024 Gazetteer, Census 2020 ZCTA-to-county file, HRSA Unmet Need Score file (ZIP-to-ZCTA crosswalk). Downloaded 2026-10-23.",
        f"Background, from HRSA's bulletin of 2026-08-27: {facts['new_access_points_2026']['millions_usd']} million dollars in New Access Points awards for {facts['new_access_points_2026']['new_sites']} new sites (announced {facts['new_access_points_2026']['announced']}).",
        "No macros, no external links. Free to use, copy and share.",
        "Code and method: https://github.com/primordial-analytics/primordial-analytics.github.io/tree/master/analysis/" + slug,
    ]
    for i, t in enumerate(lines, start=3):
        ws.cell(i, 1, t).alignment = WRAP
        if t:
            ws.row_dimensions[i].height = 18 + 15 * (len(t) // 120)

    # ---------------- Start here
    ws = start
    ws.column_dimensions["A"].width = 110
    ws["A1"] = "Site screen for health center planning teams"
    ws["A1"].font = TITLE
    intro = [
        "An example from Primordial Analytics on public data. It is not client work and uses no patient information.",
        "",
        "What it does. For any ZIP code area in the U.S., it estimates who a typical single health center site's patients would be: the share with no health center today, the share taken from neighbours, and the payer mix. You supply the patients a year.",
        "",
        "How to use it.",
        "1. Open the Screen sheet. Type a 5-digit ZIP code in column A and the patients a year the site could see in column B (yellow cells).",
        "2. Read across: the share who have no center today (low, likely, high) and, where a typical site would fill your capacity, the number of patients, those taken from neighbours and the payer mix. Where it would not, the Note column says so.",
        "3. Compare several ZIP codes by filling more rows. The ZIP table sheet lists all ZIP code areas and can be filtered and sorted.",
        "",
        "What it cannot do. It is a scenario for a typical center, not a forecast. Staffing, hours, language and referrals are not in the data. See the last sheet.",
        "Open it in Excel, LibreOffice or Numbers. A previewer that does not calculate (some mail and phone viewers) shows blank result cells.",
    ]
    for i, t in enumerate(intro, start=3):
        ws.cell(i, 1, t).alignment = WRAP
        if t:
            ws.row_dimensions[i].height = 32
    wb.calculation.fullCalcOnLoad = True
    out_dir.mkdir(exist_ok=True, parents=True)
    path = out_dir / ("check.xlsx" if for_check else f"{slug}-site-screening.xlsx")
    wb.save(path)
    add_cached_values(path)
    return path
