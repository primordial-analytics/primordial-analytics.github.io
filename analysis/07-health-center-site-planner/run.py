"""Entry point: python run.py

Downloads the public files, fits the model, runs the checks, and writes output/ (numbers.json holds every number quoted on the
page), the downloads in downloads/ and the page data in assets/examples/07/. Heavy intermediate results are cached in data/raw/cache.
"""
import json
import shutil
import sys
import tempfile
import time
from pathlib import Path

sys.dont_write_bytecode = True  # no __pycache__ in the repo (PYTHONDONTWRITEBYTECODE is ignored under python -I)
sys.path.insert(0, str(Path(__file__).parent))  # also works under python -I

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

import fetch  # noqa: E402
import fits  # noqa: E402
import model  # noqa: E402
import prepare  # noqa: E402
import project  # noqa: E402
import reach  # noqa: E402
import results  # noqa: E402
import sources  # noqa: E402
import validate  # noqa: E402
import web_data  # noqa: E402
import workbook  # noqa: E402

HERE = Path(__file__).parent
OUT = HERE / "output"
SITE = HERE.parent.parent
SLUG = "07-health-center-site-planner"
T0 = time.time()


def log(msg):
    print(f"[{time.time() - T0:6.0f}s] {msg}", flush=True)


def j(x):
    """JSON-safe copy: numpy scalars and arrays to Python, NaN to None."""
    if isinstance(x, dict):
        return {str(k): j(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [j(v) for v in x]
    if isinstance(x, np.ndarray):
        return j(x.tolist())
    if isinstance(x, (np.floating, float)):
        return None if x != x else float(x)
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.bool_,)):
        return bool(x)
    return x


def rnd(x, nd=3):
    return round(float(x), nd)


def entity_names(data):
    """Display name of every modeled entity. Look-alikes have only a grant ID in the UDS file, so they carry their first site's name."""
    site = data.sites.groupby("grant").site.first()
    names = data.ent.name.to_numpy().astype(object)
    lal = data.free == 0
    names[lal] = [f"{site.get(g, g)} (look-alike)" for g in data.ent.index[lal]]
    return names


def centers_with_area(d, data, lowd, patients_total, notp, radius=40.0):
    """Per center: the ZCTAs within `radius` miles of its nearest site and the totals there (service area for the page)."""
    import scipy.sparse as sp
    m = (data.d <= radius) & (data.free[data.c] > 0)
    M = sp.csr_matrix((np.ones(m.sum()), (data.c[m], data.z[m])), shape=(len(data.n_sites), len(data.L)))
    nq = np.stack([np.percentile(M @ notp, p, axis=1) for p in (10, 50, 90)], axis=1)
    cen = d["centers"].loc[data.ent.index[data.free > 0]]
    idx = data.eidx[cen.index].to_numpy()
    lat = data.sites.groupby("grant").lat.mean().reindex(cen.index)
    lon = data.sites.groupby("grant").lon.mean().reindex(cen.index)
    small = (data.L < 1000).astype(float)
    out = pd.DataFrame({
        "grant": cen.index, "name": cen.name.to_numpy(), "state": cen.state.to_numpy(), "type": cen.type.to_numpy(), "urban_flag": cen.urban_flag.to_numpy(),
        "patients": cen.patients.to_numpy(), "sites": cen.n_sites.to_numpy(), "lat": lat.to_numpy(), "lon": lon.to_numpy(),
        "area_zctas": np.asarray(M[idx].sum(axis=1)).ravel().astype(int), "area_low_income": np.asarray(M[idx] @ data.L).ravel().round().astype(int),
        "area_patients_all_centers": np.asarray(M[idx] @ patients_total).ravel().round().astype(int),
        "area_not_p10": nq[idx, 0].round().astype(int), "area_not_p50": nq[idx, 1].round().astype(int), "area_not_p90": nq[idx, 2].round().astype(int),
    })
    return out


def pick_fixtures(d, data, zcta, table, pen50, zd, sets, nbr, scr, Ks, centers_area):
    """Two ZIPs on which the projection must read correctly, and two centers for the service-area view.
    Rule A: an urban ZIP with 5,000 to 15,000 low-income residents where exactly three centers have 100 or more patients; the one
    closest to the median size among those. Rule B: a rural ZIP with 1,500 or more low-income residents, no permanent site within
    20 miles and under 10% penetration; the largest."""
    r = d["cells"]
    per = r[r.total.fillna(0) >= 100].groupby("zcta").grant.nunique().reindex(zcta.index).fillna(0)
    ok_a = (zcta.dclass == 2) & zcta.low.between(5000, 15000) & (per == 3)
    ok_b = (zcta.dclass == 0) & (zcta.low >= 1500) & (table.nearest_site_miles >= 20) & (pen50 < 0.10)
    la = zcta.low[ok_a]
    picks = {"served_by_three_centers": la.index[int(np.argmin((la - la.median()).abs().to_numpy()))], "no_site_within_20_miles": zcta.low.where(ok_b).idxmax()}
    out = {"rules": {"served_by_three_centers": "urban ZIP, 5,000 to 15,000 low-income residents, exactly three centers with 100+ patients; closest to the median size",
                     "no_site_within_20_miles": "rural ZIP, 1,500+ low-income residents, no permanent site within 20 miles, penetration under 10%; the largest"},
           "qualifying_zctas": {"served_by_three_centers": int(ok_a.sum()), "no_site_within_20_miles": int(ok_b.sum())}}
    for name, zc in picks.items():
        i = zcta.index.get_loc(zc)
        out[name] = {f"capacity_{K}": candidate_summary(data, zcta, zc, i, zd, sets, nbr, scr, K, table) for K in Ks}
    big_u = centers_area[centers_area.urban_flag == "Urban"].nlargest(1, "patients").iloc[0]
    rural = centers_area[centers_area.urban_flag == "Rural"]
    med = rural.iloc[int(np.argmin((rural.patients - rural.patients.median()).abs().to_numpy()))]
    out["centers"] = {"large_urban": j(big_u.to_dict()), "rural_median": j(med.to_dict())}
    return out


def candidate_summary(data, zcta, zc, i, zd, sets, nbr, scr, K, table):
    z, dd, res = project.candidate(zd.lat[i], zd.lon[i], zd, sets, 0, K, nbr, None)
    D = np.array([r["D"] for r in res])
    N = np.array([r["new"] for r in res])
    share = N / np.maximum(D, 1e-9)
    loss = project.neighbour_loss(z, res, nbr)
    names = entity_names(data)
    top = sorted(((float(np.median(v)), int(c)) for c, v in loss.items()), reverse=True)[:5]
    pay = np.zeros((len(res), 4))
    for b, r in enumerate(res):
        pay[b] = (r["new_z"][:, None] * zd.new_payer[z] + (r["draw_z"] - r["new_z"])[:, None] * zd.obs_mix[z]).sum(axis=0)
    sh = pay / pay.sum(axis=1, keepdims=True)
    # rank among ZCTAs within 15 miles on newly reached patients per 1,000 of capacity (small and campus ZIPs left out)
    dist = model.miles_between(zd.lat[i], zd.lon[i], zd.lat, zd.lon)
    near = np.flatnonzero((dist <= 15) & (zcta.low >= reach.SMALL_LOW_INCOME).to_numpy() & (zcta.college_pct < reach.CAMPUS_COLLEGE_PCT).to_numpy())
    newK = np.minimum(scr["draw_uncapped"], K) * np.nan_to_num(scr["share_new"])
    ranks = np.array([(newK[b][near] > newK[b][i]).sum() + 1 for b in range(newK.shape[0])])
    sources_z = np.median(np.stack([r["draw_z"] for r in res]), axis=0)
    ordz = np.argsort(-sources_z)[:5]
    return {
        "zcta": zc, "state": zcta.state.iloc[i], "county": table.county.iloc[i], "density_class": table.density_class.iloc[i], "low_income_residents": int(zcta.low.iloc[i]),
        "centers_with_patients": int(table.centers_with_patients.iloc[i]), "nearest_site_miles": float(table.nearest_site_miles.iloc[i]), "capacity": K,
        "patients_p10_p50_p90": [rnd(v, 0) for v in np.percentile(D, [10, 50, 90])], "capacity_binds_in_sets_pct": rnd(100 * np.mean([r["scale"] < 1 for r in res]), 0),
        "share_new_p10_p50_p90": [rnd(v) for v in np.percentile(share, [10, 50, 90])],
        "new_patients_p50": rnd(np.median(N), 0), "from_neighbours_p50": rnd(np.median(D - N), 0),
        "payer_share_p50": {p: rnd(v) for p, v in zip(project.PAYERS, np.median(sh, axis=0))},
        "payer_share_p10": {p: rnd(v) for p, v in zip(project.PAYERS, np.percentile(sh, 10, axis=0))},
        "payer_share_p90": {p: rnd(v) for p, v in zip(project.PAYERS, np.percentile(sh, 90, axis=0))},
        "rank_within_15_miles_p10_p50_p90": [int(v) for v in np.percentile(ranks, [10, 50, 90])], "zctas_within_15_miles": int(len(near)),
        "top_neighbours": [{"name": str(names[c]), "patients_lost_p50": rnd(v, 0)} for v, c in top],
        "top_source_zips": [{"zcta": zcta.index[z[k]], "patients_p50": rnd(sources_z[k], 0)} for k in ordz],
    }


def top_list(values, zcta, small, n=50, ascending=False):
    s = pd.Series(np.asarray(values, float), index=zcta.index)[~small]
    return list(s.sort_values(ascending=ascending).index[:n])


def main():
    OUT.mkdir(exist_ok=True)
    fetch.fetch_all()
    log("raw files present")
    d = prepare.build()
    zcta, centers = d["zcta"], d["centers"]
    facts = sources.write(OUT)

    # ---------------------------------------------------------------- fit
    x_main, info_main = fits.main_fit()
    _, data, a = fits.variant("main")
    X, boot_info = fits.bootstrap(x_main)
    log(f"main fit and {len(X)} bootstrap refits ready")
    boot_conv = {"refits": len(X), "converged": int(sum(i["ok"] for i in boot_info)), "iteration_cap": fits.BOOT_ITER, "stopped_at_cap": int(sum(i["nit"] >= fits.BOOT_ITER for i in boot_info)),
                 "iterations_p10_p50_p90": [int(v) for v in np.percentile([i["nit"] for i in boot_info], results.PCT)],
                 "max_abs_gradient_p50_max": [float(np.median([i["gmax"] for i in boot_info])), float(np.max([i["gmax"] for i in boot_info]))],
                 "messages": {m: int(sum(i["message"] == m for i in boot_info)) for m in sorted({i["message"] for i in boot_info})}}
    # Do the refits sit at their optimum? Continue the first eight for as long again and see how far the decay and the new-site level move.
    cont = fits.bootstrap_jobs("main", x_main, fits.N_BOOT, model.SEED)[:8]
    for jb in cont:
        jb["x0"], jb["tag"] = X[jb["tag"]], f"cont{jb['tag']}"
    rc = fits.run_jobs(cont, workers=8, name="boot_continue")
    pt = lambda x: model.params_table(x)
    single = (data.n_sites == 1) & (data.free > 0) & (data.type == 0)

    def th(x):  # a new single site's level, as project.Sets sets it
        q = model.unpack(x)
        return np.exp(q["mu"][0] - q["beta"] * data.lnn_ref + np.exp(q["log_sigma"][0]) * np.median(q["eps"][single]))

    th0 = lambda x: np.exp(model.unpack(x)["mu"][0] - model.unpack(x)["beta"] * data.lnn_ref)
    boot_conv["continued_refits"] = {"extra_iterations": fits.BOOT_ITER, "refits": len(cont),
                                     "max_relative_change_rho": rnd(max(np.abs(np.array(pt(rc[f"cont{b}"][0])["rho"]) / np.array(pt(X[b])["rho"]) - 1).max() for b in range(len(cont))), 4),
                                     "max_relative_change_gamma": rnd(max(np.abs(np.array(pt(rc[f"cont{b}"][0])["gamma"]) / np.array(pt(X[b])["gamma"]) - 1).max() for b in range(len(cont))), 4),
                                     "max_relative_change_new_site_level_chc": rnd(max(abs(th(rc[f"cont{b}"][0]) / th(X[b]) - 1) for b in range(len(cont))), 4),
                                     "max_relative_change_level_with_center_effect_zero_chc": rnd(max(abs(th0(rc[f"cont{b}"][0]) / th0(X[b]) - 1) for b in range(len(cont))), 4),
                                     "note": "The new-site level is what the projection uses and is pinned down; the level at a zero center effect moves along a ridge between the type level, the size trend and the spread of center effects, which the projection does not use.",
                                     "max_objective_change_loglik_units": rnd(max(abs(rc[f"cont{b}"][1]["fun"] - boot_info[b]["fun"]) for b in range(len(cont))) * model.SCALE, 1)}
    pm_ = model.unpack(x_main)
    bins_ = validate.sites_bin(data.n_sites)
    eps_by_bin = {}
    for t_, tname in enumerate(model.TYPES):
        for s_ in validate.SITE_BINS:
            m_ = (data.free > 0) & (data.type == t_) & (bins_ == s_)
            if m_.any():
                eps_by_bin.setdefault(tname, {})[s_] = {"centers": int(m_.sum()), "median_center_effect": rnd(np.median(pm_["eps"][m_]), 3)}
    lam_main, S_main = model.expected(x_main, a)
    ptab = model.params_table(x_main)
    pb = [model.params_table(x) for x in X]
    boot_spread = {k: {cls: [rnd(np.percentile([p[k][i] for p in pb], q), 4) for q in (10, 50, 90)] for i, cls in enumerate(results.CLASS_NAMES)} for k in ("rho", "gamma")}

    # ---------------------------------------------------------------- reach
    nls = float((centers.low_share * centers.patients).sum() / centers.patients[centers.low_share.notna()].sum())
    chk = {"A": validate.check_hidden_reconstruction(d, data, x_main)}
    chk["A"]["note"] = ("Censored cells hold hidden_patients_pct_of_center_totals of patients. Where they are placed changes ZIP counts, not the national total. "
                        "The fitted expectation places small cells worse than an equal split, so a flat prior with the exact center sums is used; the interval for one placed cell is nearly the whole 1 to 15 range.")
    log(f"check A done; censored cells placed with {chk['A']['placement_used']}")
    hid, y, bad = reach.place_hidden(d, data, x_main, prior=chk["A"]["placement_used"])
    patients = reach.patients_by_zcta(d, hid, y, nls)
    payer = reach.payer_counts(d, hid, y)
    mix, kappa, gmix = reach.shrunk_payer_mix(zcta, payer)
    L_draws, U_draws = reach.acs_draws(zcta)
    rt, pen, notp, lowd = reach.reach_tables(d, patients, L_draws, U_draws)
    agg = results.national(zcta, L_draws, U_draws, lowd, pen)
    sens = results.under_threshold_sensitivity(zcta, patients, L_draws, lowd)
    log("reach accounting done")
    # sensitivity of the headline to the assumed ACS error correlation between ZIPs (0 = Census approximation, 1 = fully shared)
    corr_sens = {}
    for rho_ in (0.0, 1.0):
        Ld, Ud = reach.acs_draws(zcta, rho=rho_, seed=8)
        pn = lowd[:, np.arange(Ld.shape[1]) % lowd.shape[1]] / np.maximum(Ld, 1)
        corr_sens[str(rho_)] = {"under_25pct_residents": results.q3((Ld * (pn < 0.25)).sum(0)), "not_patients": results.q3(np.clip(Ld - lowd[:, np.arange(Ld.shape[1]) % lowd.shape[1]], 0, None).sum(0))}
    hidden_info = {
        "hidden_patients": int(centers.hidden_patients.sum()), "hidden_cells": int(centers.hidden_cells.sum()),
        "hidden_patients_pct_of_center_totals": rnd(100 * centers.hidden_patients.sum() / centers.patients.sum(), 2),
        "centers_with_hidden_cells": int((centers.hidden_cells > 0).sum()), "centers_where_hidden_sum_is_impossible": len(bad),
        "draws_per_center": y.shape[1], "mean_placed_per_hidden_cell": rnd(y.mean(), 2),
    }
    vis_z = patients["total"].mean(axis=1) - patients["hidden_mean"]
    has_h = patients["hidden_mean"] > 0
    hidden_info["zctas_with_hidden_cells"] = int(has_h.sum())
    hidden_info["median_zcta_total_change_pct_where_hidden"] = rnd(100 * np.median(patients["hidden_mean"][has_h] / np.maximum(vis_z[has_h], 1)), 1)
    small_z = has_h & (vis_z < 100)
    hidden_info["zctas_with_hidden_cells_and_under_100_visible"] = int(small_z.sum())
    hidden_info["median_change_pct_zctas_under_100_visible"] = rnd(100 * np.median(patients["hidden_mean"][small_z] / np.maximum(vis_z[small_z], 1)), 1)
    hidden_info["note"] = ("Each center's hidden sum is fixed by Table 4, so placing the censored cells cannot change the national total. "
                           "They hold hidden_patients_pct_of_center_totals of patients; where they go changes ZIP counts.")
    hidden_info["payer_dirichlet_kappa_patients"] = rnd(kappa, 1)

    g_ = data.free[data.c] > 0
    calib = {name: {"observed": int(data.lo[g_ & (data.k == k)].sum()), "fitted": int(lam_main[g_ & (data.k == k)].sum())} for k, name in enumerate(results.CLASS_NAMES)}
    calib["by_center_type"] = {t: {"observed": int(data.lo[g_ & (data.type[data.c] == i)].sum()), "fitted": int(lam_main[g_ & (data.type[data.c] == i)].sum())} for i, t in enumerate(model.TYPES)}
    t4 = centers[["pay_uninsured", "pay_medicaid", "pay_medicare", "pay_private"]].sum()
    payer_check = {"table4_national_shares": {p: rnd(v, 4) for p, v in zip(project.PAYERS, t4 / t4.sum())},
                   "placed_zip_rows_shares": {p: rnd(v, 4) for p, v in zip(project.PAYERS, payer.sum(0) / payer.sum())},
                   "max_abs_difference": rnd(np.abs(t4.to_numpy() / t4.sum() - payer.sum(0) / payer.sum()).max(), 4)}

    # ---------------------------------------------------------------- projection
    sets = project.Sets(data, X, L_draws)
    new_payer = project.new_patient_payer(zcta, mix, float(payer[:, 0].sum() / payer.sum()))
    zd = project.ZData(zcta, new_payer, mix)
    zp = project.zpairs(zcta)
    log(f"projection pairs: {len(zp[0]):,}")
    sets24 = sets.subset(web_data.N_SETS)  # the ZIP table, the workbook and the page all use these 24 parameter sets, so they agree
    scr = results.screening(zp, zd, sets24, 0)
    log("screening table computed")
    table = results.zip_table(d, data, rt, mix, scr, None)
    nbr = project.top_neighbours(lam_main, data)
    covered = float(np.minimum(np.bincount(data.z, lam_main, minlength=len(zcta)), 1e18).sum())
    cover_top = float(nbr["lam"].sum() / covered)
    site_cap = (centers.patients / centers.n_sites).dropna()
    cap_q = [int(round(v)) for v in site_cap.quantile([0.25, 0.5, 0.75])]
    area = centers_with_area(d, data, lowd, patients["total"].mean(axis=1), notp)
    fixtures = pick_fixtures(d, data, zcta, table, rt.pen_p50, zd, sets, nbr, scr, [cap_q[1], results.K_SCREEN], area)
    log("fixtures done")

    # ---------------------------------------------------------------- checks that need refits
    from scipy.stats import spearmanr
    jobs, test_mask = validate.holdout_jobs(data, x_main)
    folds, finfo = validate.loco_plan(data, d)
    jobs += validate.loco_jobs(data, folds, x_main)
    for v in ("r40", "r100", "popbase"):
        jobs.append({"variant": v, "tag": v, "x0": x_main})
    jobs.append({"variant": "late", "tag": "late"})
    no_cens = np.ones(len(data.c))
    no_cens[data.censored] = 0.0
    jobs.append({"variant": "main", "tag": "nocens", "pm": no_cens, "x0": x_main})
    jobs.append({"variant": "y2024", "tag": "y2024"})
    res = fits.run_jobs(jobs, workers=8, name="validation")
    log("validation refits done")
    chk["B"], sfca_r = validate.check_holdout(d, data, a, res["holdout"][0], test_mask)
    chk["B"]["note"] = ("The negative-binomial 80% intervals cover well above 80% of held-out cells, so they are wider than needed, not calibrated. Poisson-only intervals cover far less. "
                        "The fit uses a Poisson likelihood, so intervals on new-site numbers reflect parameter uncertainty only; interval_80_coverage_poisson_only shows how much the Poisson noise model understates spread.")
    chk["C"], loco_rows = validate.check_loco(d, data, a, folds, res, finfo, sfca_r)
    log("B, C done")

    d24, data24, a24 = fits.variant("y2024")
    x24 = res["y2024"][0]
    jobs24 = fits.bootstrap_jobs("y2024", x24, 40, 5)
    res24 = fits.run_jobs(jobs24, workers=8, name="boot24")
    boots24 = np.stack([res24[b][0] for b in range(40)])
    new_sites = validate.opening_window(d24)
    chk["D"], d_rows = validate.check_openings(d24, d, x24, boots24, data24, a24, new_sites)
    log("D done")
    cmp = pd.DataFrame({"low": zcta.low, "pop": zcta["pop"], "state": zcta.state, "patients_total": rt.patients_total, "pen_p50": rt.pen_p50, "notpat_p50": rt.notpat_p50})
    chk["E"] = validate.check_hrsa(zcta, d["uns"], cmp)

    # F: sensitivity of the top-50 lists and of the share of a site's draw that is new
    small = (rt.small | rt.campus).to_numpy()  # left out of ranked lists: small or campus ZIPs
    main_people = top_list(rt.notpat_p50, zcta, small)
    unscaled = top_list(np.clip(zcta.low - rt.patients_total, 0, None), zcta, small)
    popbase = top_list(np.clip(zcta["pop"] - rt.patients_total, 0, None), zcta, small)
    k_rank = cap_q[0]  # 25th percentile site size: a capacity at which fewer ZIPs tie than at 5,000

    def screen_top(x, variant):
        _, dat, aa = fits.variant(variant)
        base = zcta["pop"] if variant == "popbase" else zcta.low
        s1 = project.Sets(dat, [x], np.repeat(base.to_numpy(np.float32)[:, None], 1, axis=1))
        sc1 = results.screening(zp, zd, s1, 0, K=k_rank)
        return sc1

    def new_at_k(sc):
        return np.minimum(sc["draw_uncapped"], k_rank) * np.nan_to_num(sc["share_new"])

    def share_stats(sc):
        return [rnd(v, 3) for v in np.percentile(np.median(np.nan_to_num(sc["share_new"]), axis=0)[~small], results.PCT)]

    def ties(v):
        s = np.sort(v[~small])[::-1]
        return int((s >= 0.99 * s[49]).sum())

    sc_main = screen_top(x_main, "main")
    main_screen = top_list(new_at_k(sc_main)[0], zcta, small)
    F = {"top50_not_patients": {"main": 50, "without_income_scaling_of_patients": validate.top_overlap(main_people, unscaled),
                                 "total_population_base": validate.top_overlap(main_people, popbase)},
         "top50_new_patients_at_capacity": {"capacity": k_rank, "main": 50, "zips_within_1pct_of_the_50th_value": ties(new_at_k(sc_main)[0])},
         "share_new_by_zip_p10_p50_p90": {"main_1_set": share_stats(sc_main), "main_page_sets": share_stats(scr)}}
    for v, label in (("r40", "pairs_within_40_miles"), ("r100", "pairs_within_100_miles"), ("popbase", "total_population_base"), ("late", "sites_added_in_2026_included")):
        sc_v = screen_top(res[v][0], v)
        F["top50_new_patients_at_capacity"][label] = validate.top_overlap(main_screen, top_list(new_at_k(sc_v)[0], zcta, small))
        F["share_new_by_zip_p10_p50_p90"][label] = share_stats(sc_v)
    F["params"] = {v: model.params_table(res[v][0]) for v in ("r40", "r100", "popbase", "late", "nocens")}
    F["new_patients_top50_main_vs_bootstrap_median_set"] = validate.top_overlap(main_screen, top_list(np.median(new_at_k(scr), axis=0), zcta, small))
    chk["F"] = F
    log("E, F done")

    # ---------------------------------------------------------------- 2SFCA agreement with the gravity model
    sp = validate.site_pairs(data, 30.0)
    train_patients = np.bincount(data.c, data.lo * (data.free[data.c] > 0), minlength=len(data.n_sites))
    _, access = validate.sfca_pairs(data, sp, train_patients, sfca_r)
    okz = ~small & (access > 0)
    low_access = pd.Series(access[okz]).rank(pct=True).to_numpy() <= 0.1
    low_pen = pd.Series(rt.pen_p50.to_numpy()[okz]).rank(pct=True).to_numpy() <= 0.1
    chk["sfca_vs_gravity"] = {
        "radius_miles": sfca_r, "spearman_access_vs_penetration": rnd(spearmanr(access[okz], rt.pen_p50[okz])[0]),
        "spearman_access_vs_not_patients_rank": rnd(spearmanr(access[okz], -rt.notpat_p50[okz])[0]),
        "lowest_tenth_by_access_also_lowest_tenth_by_penetration_pct": rnd(100 * (low_access & low_pen).sum() / low_access.sum(), 1),
        "cannot_answer": ["capacity of a new site", "competition among sites", "where a new site's patients come from", "payer mix of the draw"]}

    # ---------------------------------------------------------------- downloads and page data
    down = SITE / "downloads"
    down.mkdir(exist_ok=True)
    loco_rows.round(4).to_csv(OUT / "check_c_leave_one_center_out.csv", index=False)
    d_rows.assign(added=d_rows.added.dt.date).round(2).to_csv(OUT / "check_d_openings.csv", index=False)
    csv_cols = list(table.columns)
    table.to_csv(OUT / "zip_table.csv", index_label="zcta")
    code = codebook(["zcta"] + csv_cols)
    code.to_csv(OUT / "zip_table_codebook.csv", index=False)
    shutil.copy(OUT / "zip_table.csv", down / f"{SLUG}-zip-table.csv")
    shutil.copy(OUT / "zip_table_codebook.csv", down / f"{SLUG}-zip-table-codebook.csv")
    ex = fixtures["served_by_three_centers"][f"capacity_{results.K_SCREEN}"]
    workbook.make(table, {"zcta": ex["zcta"], "capacity": results.K_SCREEN}, facts, down, SLUG, cap_q[1], cap_q[0], cap_q[2], results.K_SCREEN)
    log("downloads written")

    manifest, blob, lists = web_data.encode(d, data, rt, table, zd, sets24, nbr, S_main)
    meta = {"sets": [{"rho": [rnd(v, 5) for v in sets24.rho[b]], "gamma": [rnd(v, 5) for v in sets24.gamma[b]], "theta": [rnd(v, 5) for v in sets24.theta[b]]} for b in range(web_data.N_SETS)],
            "default_capacity": cap_q[1], "capacity_p25_p75": [cap_q[0], cap_q[2]], "radius_miles": model.RADIUS, "area_radius_miles": 40,
            "density_cuts_per_sq_mile": list(prepare.DENSITY_CUTS), "centers": j(area.to_dict("records")), "facts": facts,
            "type_labels": {"chc": "Health center", "special": "Migrant, homeless or public housing funding only"},
            "lookalikes": [{"name": str(n), "state": str(s)} for n, s in zip(entity_names(data)[data.free == 0], data.ent.state[data.free == 0])],
            "note": "Capacity is an input, not an output: the model does not forecast how many patients a site draws, it says who they would be."}
    gz_bytes, json_bytes = web_data.write(SITE / "assets" / "examples" / "07", manifest, blob, lists, meta)
    browser_check = web_data.check(SITE / "assets" / "examples" / "07", zd, sets, nbr)
    wb_check = check_workbook(table, facts, zcta, SITE / "assets" / "examples" / "07")
    top20 = center_top_zips(d, zcta)
    (SITE / "assets" / "examples" / "07" / "centers_top_zips.json").write_text(json.dumps(top20, separators=(",", ":")))
    log("page data written")

    # ---------------------------------------------------------------- numbers.json
    top_tab = table[(table.small_zip == 0) & (table.campus_zip == 0)].nsmallest(50, "rank_people")[["state", "county", "low_income_residents", "patients_all_centers", "penetration_p50", "not_patients_p10", "not_patients_p50", "not_patients_p90",
                                                                         "hrsa_published_penetration", "ratio_to_hrsa", "hrsa_disagrees", "no_patients_reported"]]
    c = chk["C"]
    oz = d["rows"][d["rows"].kind == "Other ZIP Codes"].groupby("grant").total.sum()
    oz_share = (oz / centers.patients.reindex(oz.index)).dropna()
    st = d["site_stats"]
    numbers = {
        "built": time.strftime("%Y-%m-%d"), "data_downloaded": (HERE / "data" / "raw" / "SOURCES.txt").read_text().splitlines()[0].replace("Downloaded ", ""),
        "data_checks": {**j(d["summary"]), "sites": j(d["site_stats"]), "model_pairs": int(len(data.c)), "model_centers": int(data.free.sum()), "lookalike_competitors": int((data.free == 0).sum()),
                        "cells_outside_60_miles": int(len(data.cells_outside)), "patients_outside_60_miles": int(data.cells_outside.V.sum()),
                        "censored_pairs": int(len(data.censored)), "late_sites_excluded_from_main": int(len(d["late_sites"]))},
        "model": {"main": ptab, "fit": info_main, "bootstrap_sets": len(X), "bootstrap_mean_iterations": rnd(np.mean([i["nit"] for i in boot_info]), 0),
                  "bootstrap_spread_p10_p50_p90": boot_spread, "density_cuts_per_sq_mile": list(prepare.DENSITY_CUTS),
                  "typical_theta_new_site": {"chc": rnd(sets.theta[:, 0].mean(), 3), "special": rnd(sets.theta[:, 1].mean(), 3)},
                  "typical_theta_if_center_effect_zero": {t_: rnd(np.mean([np.exp(model.unpack(x)["mu"][i] - model.unpack(x)["beta"] * data.lnn_ref) for x in X]), 3) for i, t_ in enumerate(model.TYPES)},
                  "center_effect_median_by_type_and_sites": eps_by_bin, "bootstrap_convergence": boot_conv,
                  "main_fit_max_abs_gradient": info_main["gmax"],
                  "low_income_share_of_patients_used_for_missing": rnd(nls, 3)},
        "hidden_cells": hidden_info, "in_sample_calibration": calib, "payer_fill_check": payer_check,
        "reach": agg, "reach_acs_correlation_sensitivity": corr_sens, "reach_sensitivity": sens,
        "projection": {"capacity_for_screening": results.K_SCREEN, "site_size_patients_per_site_p25_p50_p75": cap_q,
                       "capacity_is_an_input": "The model gives who a site's patients would be, not how many. Patients = the capacity chosen.",
                       "neighbour_lists_cover_share_of_expected_patients": rnd(cover_top, 3), "fixtures": fixtures,
                       "share_new_p50_median_zip": rnd(table.site_share_new_p50[(table.small_zip == 0) & (table.campus_zip == 0)].median(), 3),
                       "share_from_neighbours_median_case": rnd(1 - table.site_share_new_p50[(table.small_zip == 0) & (table.campus_zip == 0)].median(), 3),
                       "zips_where_site_would_be_full_pct": rnd(100 * (np.median(scr["draw_uncapped"], axis=0)[~small] > results.K_SCREEN).mean(), 1),
                       "zips_where_capacity_binds_pct": rnd(100 * table.site_capacity_binds[(table.small_zip == 0) & (table.campus_zip == 0)].mean(), 1)},
        "checks": chk,
        "tool_checks": {"browser_formula_vs_python": browser_check, "workbook_vs_python": wb_check, "page_data_bytes_gz": gz_bytes, "page_data_json_bytes": json_bytes},
        "top50_zips_by_people_not_patients": j(top_tab.reset_index(names="zcta").to_dict("records")),
        "facts_file": "output/facts.json",
        "definitions": {
            "low_income": "under 200% of the federal poverty level (ACS table C17002, estimates 002 to 007)",
            "patients_low_income": "a center's patients in a ZIP times that center's share at or below 200% of poverty (Table 4 lines 1 to 3 over lines 1 to 4, unknown excluded); assumes a center's ZIP mix equals its overall mix",
            "penetration": "low-income patients of all centers over low-income residents; ours, not HRSA's (HRSA's published value is a comparator)",
            "not_patients": "low-income residents minus low-income patients, floored at 0; one measure, they may use other care",
            "under_10_and_25_pct": "our own thresholds on penetration, named as ours on the page; not an HRSA rule",
            "small_zip": "under 1,000 low-income residents: ACS margins are wide; left out of ranks and ranked tables, counted in totals",
            "campus_zip": "college and graduate students (ACS B14001) are 20% or more of residents aged 3 and over. The poverty count includes students living off campus, so these ZIPs look poorer and less reached than their non-student residents; left out of ranks and ranked tables, counted in totals",
            "typical_site": "a single new site with the median attractiveness of its type (center type effect at one site)",
            "draw": "model quantity behind patients = min(capacity, draw). Not published: most single sites would exceed any plausible capacity, so capacity is an input and the projection gives shares (new, from neighbours, payer mix)",
            "interval": "reach figures: 10th to 90th percentile over Monte Carlo draws of the ACS margins and of the censored cells (2,000 draws). New-site shares: over 24 of the 100 bootstrap refits of the model. Survey and parameter uncertainty only: yearly patient-count noise and the assumptions in 'limits' are not in it",
            "tile_3": f"mean over {c['centers']} held-out centers of the share of a center's patients living in the 20 ZIPs the model ranked highest for it; patient_weighted weights centers by size. Baselines: the 20 nearest ZIPs and the 20 ZIPs with most low-income residents",
            "geography": "50 states, DC and Puerto Rico. Virgin Islands, Guam, American Samoa and Northern Mariana patients are in Table 4 totals but in no ZIP table",
            "new_patients": "patients who, in the model, use no health center today (the outside option); the rest of a site's patients are taken from neighbouring centers",
        },
        "limits": limits_text(d, st, oz_share, data, chk, sens),
        "headline": headline(agg, c, sens, corr_sens),
        "sources": [line.split("\t") for line in (HERE / "data" / "raw" / "SOURCES.txt").read_text().splitlines()[1:]],
    }
    (OUT / "numbers.json").write_text(json.dumps(j(numbers), indent=1))
    log("numbers.json written")
    web_data.finish(SITE / "assets" / "examples" / "07", OUT)
    log("page data finished")
    web_data.check_page_text(SITE / "_casestudies" / f"{SLUG}.md", OUT)
    log("front matter and static text agree with output/")


def headline(agg, c, sens, corr):
    u, u10 = agg["under_25pct"], agg["under_10pct"]
    s = sens["patients_scaled_to_income_share"]
    return {
        "tile_1_residents_where_under_25pct_reached": {
            "p10_p50_p90": u["low_income_residents"], "interval_covers": "survey margin of error only",
            "interval_if_acs_errors_shared_within_state": corr["1.0"]["under_25pct_residents"],
            "if_patients_not_scaled_to_income_share_p10_p50_p90": sens["patients_not_scaled"]["under_25pct"]["low_income_residents"],
            "at_20pct_p50": s["under_20pct"]["low_income_residents"][1], "at_30pct_p50": s["under_30pct"]["low_income_residents"][1],
            "uninsured_adults_same_zips_p10_p50_p90": u["uninsured_adults_below_200"], "zctas": u["zctas"]},
        "tile_2_residents_where_under_10pct_reached": {"p10_p50_p90": u10["low_income_residents"], "zctas": u10["zctas"], "interval_covers": "survey margin of error only"},
        "tile_3_of_100_patients_in_model_top20": rnd(100 * c["top20_model_patient_weighted"], 1), "tile_3_mean_over_centers": rnd(100 * c["top20_model_mean"], 1),
        "tile_3_nearest_20_zips_patient_weighted": rnd(100 * c["top20_nearest_zips_patient_weighted"], 1), "tile_3_nearest_20_zips_mean": rnd(100 * c["top20_nearest_zips_mean"], 1),
        "tile_3_population_only": rnd(100 * c["top20_population_only_mean"], 1),
        "tile_3_share_of_centers_model_beats_nearest": rnd(100 * c["share_centers_model_beats_nearest"], 0), "tile_3_median_points_over_nearest": rnd(c["model_minus_nearest_points_median"], 1),
        "tile_3_single_site_centers": {k: (v if k == "centers" else rnd(v, 3)) for k, v in c["by_sites_bin"]["1"].items()},
    }


def limits_text(d, st, oz_share, data, chk, sens):
    """Plain statements of what the numbers leave out, with their sizes from the run."""
    sm = d["summary"]
    sh = chk["F"]["share_new_by_zip_p10_p50_p90"]
    gap = (sens["patients_scaled_to_income_share"]["under_25pct"]["low_income_residents"][1] - sens["patients_not_scaled"]["under_25pct"]["low_income_residents"][1]) / 1e6
    return [
        f"'Other ZIP Codes' lines hold {sm['visible_patients_other_zip']:,} patients ({100 * sm['visible_patients_other_zip'] / sm['national_patients_table4']:.1f}% of all patients, {100 * oz_share.mean():.1f}% of the average center and up to {100 * oz_share.max():.0f}% of one). They cannot be placed on a ZIP, so they are in no ZIP's count.",
        f"Look-alike centers ({int((data.free == 0).sum())} modeled as competitors) have no patients in the UDS awardee file. Their patients are not counted, so they show as people not reached.",
        "A person served by two centers is counted twice: UDS patients are unduplicated within a center only.",
        f"Supply is permanent, non-school sites. The site list has {st['rows_seasonal']:,} seasonal rows, {st['rows_mobile_van']:,} mobile van rows and {st['rows_school_setting']:,} rows in a school setting that are left out. A mobile unit or outreach team is not in the supply.",
        "ACS counts college students as low-income, which lifts the counts of some ZIPs near campuses.",
        "Distance is straight-line miles to the center's nearest permanent site. Drive time, transit and borders are not in it.",
        f"Patients' income is each center's overall low-income share applied to every ZIP of that center. Without that scaling the residents in ZIPs under 25% reached fall by {gap:.1f} million.",
        "The thresholds of 10% and 25% penetration are ours, not an HRSA rule.",
        "Sites that closed are missing from the October 2026 site list, which shifts the competition around the 2025 openings.",
        f"Intervals are parameter and survey uncertainty. Yearly patient-count noise, the capacity and who could use a center are not in them: the median share of a site's patients with no center today is {sh['main_page_sets'][1]:.2f} with low-income residents as the market and {sh['total_population_base'][1]:.2f} with all residents.",
    ]


def center_top_zips(d, zcta):
    """Each center's top 20 ZIPs by patients, for the lazy-loaded service-area list."""
    r = d["cells"]
    r = r[r.total.fillna(0) > 0]
    zi = pd.Series(np.arange(len(zcta)), index=zcta.index)
    g = r.groupby(["grant", "zcta"]).total.sum().reset_index().sort_values(["grant", "total"], ascending=[True, False]).groupby("grant").head(20)
    return {k: [[int(zi[z]), int(t)] for z, t in zip(v.zcta, v.total)] for k, v in g.groupby("grant")}


def codebook(cols):
    desc = {
        "zcta": ("ZIP code tabulation area, 5 digits as text (Excel drops leading zeros unless the column is read as text)", "Census 2024 Gazetteer; ACS 2020-2024 geography"),
        "state": ("State of the ZCTA's main county", "Census 2020 ZCTA-to-county file"),
        "county": ("Main county (largest land overlap). 'city', Borough, Census Area and Municipality names are kept, so a Virginia independent city or an Alaska borough is not merged with a neighbouring county. Connecticut has its eight legacy counties, not the planning regions", "Census 2020 ZCTA-to-county file"),
        "density_class": ("rural under 200, suburban 200 to 2,000, urban 2,000 or more residents per square mile of land", "ACS population over Gazetteer land area"),
        "residents": ("Population for whom poverty status is determined", "ACS 2020-2024 table C17002"),
        "low_income_residents": ("Residents under 200% of the federal poverty level", "ACS C17002, estimates 002 to 007"),
        "low_income_moe90": ("90% margin of error on low_income_residents", "ACS C17002 (margins summed in quadrature)"),
        "uninsured_adults_below_200": ("Uninsured adults 19 to 64 under 200% of poverty", "ACS C27016, estimates 008 + 018 + 028"),
        "centers_with_patients": ("Health centers reporting at least one patient living here", "HRSA UDS 2025, ZIP table"),
        "nearest_site_miles": ("Straight-line miles from the ZCTA centroid to the nearest permanent, non-school site", "HRSA site list, October 2026"),
        "patients_all_centers": ("Patients of all health centers living here, censored cells placed (mean)", "HRSA UDS 2025; placement in analysis/hidden.py"),
        "low_income_patients": ("Patients of all health centers living here, scaled by each center's share of patients under 200% of poverty (mean)", "HRSA UDS 2025 Table 4 lines 1 to 4"),
        "patients_per_100_residents": ("patients_all_centers per 100 residents, no income scaling. Blank where small_zip is 1 or no_patients_reported is 1", "derived"),
        "penetration_p10": ("Low-income patients over low-income residents, 10th percentile. Blank where small_zip is 1 or no_patients_reported is 1; over_100pct marks ZIPs whose patients exceed residents", "Monte Carlo on ACS margins and censored cells"),
        "penetration_p50": ("Low-income patients over low-income residents, median. Blank where small_zip is 1 or no_patients_reported is 1", "Monte Carlo"), "penetration_p90": ("Low-income patients over low-income residents, 90th percentile. Blank where small_zip is 1 or no_patients_reported is 1", "Monte Carlo"),
        "not_patients_p10": ("Low-income residents who are not health center patients, 10th percentile", "residents minus low-income patients, floored at 0"),
        "not_patients_p50": ("Low-income residents who are not health center patients, median (residents minus low-income patients, floored at 0)", "Monte Carlo"), "not_patients_p90": ("Low-income residents who are not health center patients, 90th percentile (residents minus low-income patients, floored at 0)", "Monte Carlo"),
        "hrsa_published_penetration": ("HRSA's published penetration of low-income residents (comparator; year not stated in the file). It uses the same UDS ZIP table, so it checks our joins, not the model. The correlation on the page includes small ZIPs, whose reach is blank here", "HRSA Unmet Need Score file"),
        "ratio_to_hrsa": ("penetration_p50 over hrsa_published_penetration, ZIPs that are not small. Ours runs lower by about the income share applied to patients", "derived"),
        "hrsa_disagrees": ("1 if ratio_to_hrsa is below 0.5 or above 2; read these rows with care (often a different reporting year)", "derived"),
        "over_100pct": ("1 if the median penetration is above 1: patients living in the ZIP exceed its low-income residents (ACS error, college or seasonal residents, or an HRSA ZIP merged into the ZCTA)", "derived"),
        "no_patients_reported": ("1 if no center reports a patient living here. Their patients, if any, sit in the centers' 'Other ZIP Codes' lines. Penetration is left blank; not_patients is then the low-income residents", "HRSA UDS 2025, ZIP table"),
        "small_zip": ("1 if under 1,000 low-income residents: ACS margins are wide; left out of ranks", "derived"),
        "college_students_pct": ("College and graduate students per 100 residents aged 3 and over", "ACS 2020-2024 table B14001, estimates 008 + 009 over 001"),
        "campus_zip": ("1 if college_students_pct is 20 or more: students inflate the low-income count; left out of ranks", "derived"),
        "rank_people": ("Rank by not_patients_p50, 1 = most, ZIPs that are neither small nor campus", "derived"),
        "rank_share": ("Rank by penetration_p50, 1 = lowest, ZIPs that are neither small nor campus", "derived"),
        "site_share_new_p10": (f"Share of a typical single site's patients at the ZIP centroid who have no health center today, 10th percentile. Does not depend on capacity", "model; the 24 bootstrap parameter sets the page uses"),
        "site_share_new_p50": ("Share of a typical single site's patients at the ZIP centroid who have no health center today, median. Does not depend on capacity", "model"), "site_share_new_p90": ("Share of a typical single site's patients at the ZIP centroid who have no health center today, 90th percentile. Does not depend on capacity", "model"),
        f"site_new_patients_at_{results.K_SCREEN}_p10": (f"New patients (no center today) if the site sees {results.K_SCREEN:,} patients a year: capacity times site_share_new, 10th percentile. Valid where site_capacity_binds is 1, because only there does capacity, not the model, set the number of patients. Capacity is an input; the model does not forecast volume", "model"),
        f"site_new_patients_at_{results.K_SCREEN}_p50": (f"New patients (no center today) if the site sees {results.K_SCREEN:,} patients a year, median. Valid where site_capacity_binds is 1", "model"), f"site_new_patients_at_{results.K_SCREEN}_p90": (f"New patients (no center today) if the site sees {results.K_SCREEN:,} patients a year, 90th percentile. Valid where site_capacity_binds is 1", "model"),
        "site_capacity_binds": (f"1 if the model's draw for a typical single site here exceeds {results.K_SCREEN:,} patients a year in 90% of parameter sets, so capacity (not the model) sets the volume. The draw itself is not published", "model"),
        "site_payer_uninsured": ("Share of the site's patients uninsured (median)", "model; ACS C27016 for new patients, observed mix for neighbours' patients"),
        "site_payer_medicaid": ("Share on Medicaid, CHIP or other public insurance (median)", "model"),
        "site_payer_medicare": ("Share on Medicare (median)", "model"), "site_payer_private": ("Share with private insurance (median)", "model"),
    }
    return pd.DataFrame([{"column": c, "meaning": desc[c][0], "source": desc[c][1]} for c in cols])


def check_workbook(table, facts, zcta, tool_dir, n=60, seed=9):
    """Fill a copy of the workbook with random ZIPs and capacities, recalculate it in LibreOffice and compare it with (a) the same
    formula in Python on the ZIP table and (b) the page's formula on the shipped 24 parameter sets."""
    rng = np.random.default_rng(seed)
    pick = rng.choice(len(zcta), n, replace=False)
    rows = [(zcta.index[i], float(rng.choice([800, 2000, 5000, 12000]))) for i in pick]
    meta, cols = web_data.decode(tool_dir)
    s_dec, zd_dec, nbr_dec = web_data.decoded_sets(meta, cols)
    with tempfile.TemporaryDirectory() as tmp:
        path = workbook.make(table, {"zcta": rows[0][0], "capacity": rows[0][1]}, facts, Path(tmp), SLUG, 0, 0, 0, results.K_SCREEN, for_check=rows)
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True)["Screen"]
        same, tool, shown, blank_ok = [], [], 0, True
        for k, (zc, cap) in enumerate(rows):
            r = workbook.FIRST + k
            t = table.loc[zc]
            new_got, share_got = ws[f"N{r}"].value, ws[f"K{r}"].value
            counts_valid = cap <= results.K_SCREEN and t.site_capacity_binds == 1
            same.append(abs(share_got - t.site_share_new_p50))
            if not counts_valid:  # the workbook shows shares only: a typical site here may not fill this capacity
                blank_ok &= all(ws[f"{c}{r}"].value in (None, "") for c in "MNOPQRST")
                continue
            shown += 1
            i = zcta.index.get_loc(zc)
            _, _, res = project.candidate(zd_dec.lat[i], zd_dec.lon[i], zd_dec, s_dec, 0, cap, nbr_dec, None)
            page_new = np.median([x["new"] for x in res])  # the page's number: min(draw, capacity) times the new share, median over the sets
            tool.append((abs(new_got - page_new) / max(page_new, 50), abs(new_got - cap * t.site_share_new_p50)))
    same, tool = np.array(same), np.array(tool)
    return {"rows": n, "rows_with_counts": shown, "counts_blank_where_capacity_may_not_bind": bool(blank_ok),
            "workbook_vs_python_on_zip_table": {"max_abs_diff_share_new": float(same.max()), "max_abs_diff_new_patients": float(tool[:, 1].max())},
            "workbook_vs_page_formula": {"max_relative_diff_new_patients": rnd(tool[:, 0].max(), 4), "median_relative_diff_new_patients": rnd(np.median(tool[:, 0]), 4)},
            "note": "Workbook formulas recalculated in LibreOffice. Patient counts are shown only where the capacity is 5,000 or less and a typical site there draws more than 5,000 (site_capacity_binds = 1), where they equal the page's min(draw, capacity) times the new share; elsewhere only the shares show. The workbook, the ZIP table and the page use the same 24 of the 100 parameter sets; the remaining difference is the page data's rounding."}


if __name__ == "__main__":
    main()
