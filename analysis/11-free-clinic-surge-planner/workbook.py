"""The downloadable planner: live formulas over the precomputed county quantiles, no macros."""
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import model
import results

NAVY, GOLD = "191345", "B3A369"
HEAD = PatternFill("solid", fgColor=GOLD)
INPUT = PatternFill("solid", fgColor="FFF2B3")
CALC = PatternFill("solid", fgColor="EFEFEF")
BOLD = Font(bold=True, color=NAVY)
TITLE = Font(bold=True, size=16, color=NAVY)
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

MONTHLY_EXAMPLE = [112, 118, 121, 115, 126, 119, 123, 127, 118, 125, 121, 115]


def header(ws, row, text, cols=6):
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = HEAD
    ws.cell(row, 1, text).font = BOLD


def put(ws, ref, value, fmt=None, fill=None, bold=False):
    cell = ws[ref]
    cell.value = value
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
        cell.border = BOX
    if bold:
        cell.font = BOLD
    return cell


def counties_sheet(wb, table):
    ws = wb.create_sheet("Counties")
    cols = ["county_label"] + [c for c in table.columns if c != "state_name"]
    ws.append(cols)
    for _, row in table.iterrows():
        label = f"{row['county']}, {row['state']}"
        vals = [label] + [None if (isinstance(row[c], float) and row[c] != row[c]) else row[c] for c in cols[1:]]
        ws.append(vals)
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).font = BOLD
        ws.cell(1, c).fill = HEAD
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(table) + 1}"
    return len(table) + 1, len(cols), cols


def add_cached_values(path):
    """openpyxl writes formulas without results, so a viewer that does not calculate shows blank cells. Recalculate a
    copy in LibreOffice (headless) and write its results into the original file, which keeps our formatting, panes
    and validation. Skipped with a message if LibreOffice is not installed."""
    soffice = shutil.which("soffice")
    if not soffice:
        print("soffice not found: workbook saved without cached values")
        return
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)], check=True, capture_output=True)
        calc = load_workbook(Path(tmp) / path.name, data_only=True)
    names = load_workbook(path, read_only=True).sheetnames
    cell = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: [a-z]+="[^"]*")*)><f>(.*?)</f><v ?/></c>', re.S)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    for i, name in enumerate(names, start=1):
        if name == "Counties":
            continue
        ws = calc[name]
        def fill(m):
            v = ws[m.group(1)].value
            if v is None:
                return m.group(0)
            if isinstance(v, str):
                return f'<c r="{m.group(1)}"{m.group(2)} t="str"><f>{m.group(3)}</f><v>{escape(v)}</v></c>'
            return f'<c r="{m.group(1)}"{m.group(2)}><f>{m.group(3)}</f><v>{v!r}</v></c>'
        key = f"xl/worksheets/sheet{i}.xml"
        members[key] = cell.sub(fill, members[key].decode()).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in members.items():
            z.writestr(n, data)


def make(table, ex_block, out_dir):
    out_dir = Path(out_dir)
    ex = results.EXAMPLE
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    plan = wb.create_sheet("Planner")
    need = wb.create_sheet("Need statement")
    assume = wb.create_sheet("Assumptions")
    intake = wb.create_sheet("Intake guide")
    last_row, ncols, cols = counties_sheet(wb, table)
    rng = f"Counties!$A$1:${get_column_letter(ncols)}${last_row}"
    hdr = f"Counties!$1:$1"
    label_rng = f"Counties!$A$2:$A${last_row}"
    default_label = ex_block['county']

    # ---------------- Planner
    ws = plan
    ws.column_dimensions["A"].width = 62
    for col in "BCDE":
        ws.column_dimensions[col].width = 17
    ws.column_dimensions["G"].width = 14
    ws["A1"] = "Free Clinic Demand Planner"
    ws["A1"].font = TITLE
    ws["A2"] = "Yellow cells are yours to change. The numbers in them now describe a hypothetical clinic, not a real one."
    header(ws, 4, "1. Your county (HealthCare.gov states only)")
    put(ws, "A5", "County (the one shown is an example default: pick yours)")
    put(ws, "B5", default_label, fill=INPUT)
    ws.merge_cells("B5:D5")
    put(ws, "A6", "Uninsured adults 18-64 at or below 200% of poverty, 2024 (Census SAHIE)")
    put(ws, "A7", "Margin of error on that estimate (90%, plus or minus)")
    put(ws, "A8", "State expanded Medicaid?")
    put(ws, "G4", "Helpers", bold=True)
    put(ws, "G5", f'=MATCH(B5,{label_rng},0)+1', fill=CALC)

    def look(name_expr, row="$G$5"):
        return f"INDEX({rng},{row},MATCH({name_expr},{hdr},0))"

    put(ws, "B6", f'={look(chr(34) + "uninsured_adults_le200_2024" + chr(34))}', "#,##0", CALC)
    put(ws, "B7", f'={look(chr(34) + "uninsured_moe90" + chr(34))}', "#,##0", CALC)
    put(ws, "B8", f'={look(chr(34) + "medicaid_expansion_state" + chr(34))}', None, CALC)
    header(ws, 10, "2. Your clinic")
    put(ws, "A11", "Unique uninsured adult patients, last 12 months")
    put(ws, "B11", ex["patients"], "#,##0", INPUT)
    put(ws, "A12", "Visits by uninsured adults, each of the last 12 months (oldest first; blanks are skipped)", bold=True)
    for i, v in enumerate(MONTHLY_EXAMPLE):
        put(ws, f"A{13 + i}", f"   Month {i + 1}")
        put(ws, f"B{13 + i}", v, "#,##0", INPUT)
    put(ws, "A25", "Visits a year (average month x 12)", bold=True)
    put(ws, "B25", "=AVERAGE(B13:B24)*12", "#,##0", CALC)
    put(ws, "A26", "Patients seen in one session")
    put(ws, "B26", ex["patients_per_session"], "0", INPUT)
    put(ws, "A27", "Weeks the clinic runs in a year")
    put(ws, "B27", 50, "0", INPUT)
    put(ws, "A28", "Volunteers on one session")
    put(ws, "B28", ex["volunteers_per_session"], "0", INPUT)
    put(ws, "A29", "Hours in one volunteer shift")
    put(ws, "B29", ex["hours_per_shift"], "0.0", INPUT)
    put(ws, "A30", "Shifts one volunteer works in a year")
    put(ws, "B30", ex["shifts_per_volunteer"], "0", INPUT)
    put(ws, "A31", "Medication cost per visit ($)")
    put(ws, "B31", ex["medication_per_visit"], "$#,##0.00", INPUT)
    header(ws, 33, "3. Assumptions")
    put(ws, "A34", "Scenario")
    put(ws, "B34", "Central", fill=INPUT)
    put(ws, "A35", "Include Medicaid work requirements? (placeholder layer, expansion states)")
    put(ws, "B35", "Yes", fill=INPUT)
    put(ws, "A36", "New arrivals relative to your current rate (1 = same rate per uninsured adult)")
    put(ws, "B36", ex["rho"], "0.00", INPUT)
    put(ws, "C34", "Lower demand, Central or Higher demand", None)
    put(ws, "G6", '=IF($B$34="Lower demand","lower",IF($B$34="Higher demand","higher","central"))', fill=CALC)
    put(ws, "G7", '=IF($B$35="Yes","all","mkt")', fill=CALC)
    put(ws, "G8", "=$B$36", fill=CALC)
    header(ws, 38, "4. Result: what to plan for in 2027")
    for col, lab in zip("BCDE", ["Today", "Likely (50%)", "80% sure", "95% sure"]):
        put(ws, f"{col}39", lab, bold=True)
    labels = ["Growth in the county's uninsured pool since 2024", "Growth in your visits (only the part not already in your last 12 months)", "Visits a year", "Sessions a year",
              "Sessions a week", "Extra sessions compared with today", "Volunteer shifts a year", "Volunteers needed",
              "Volunteer hours", "Medication dollars", "New patients expected"]
    for i, lab in enumerate(labels):
        put(ws, f"A{40 + i}", lab)
    for col, q in zip("CDE", ["p50", "p80", "p95"]):
        g = look(f'"g_"&$G$6&"_"&$G$7&"_{q}"')
        put(ws, f"{col}40", f"={g}", "0.0%", CALC)
        cg = look(f'"cg_"&$G$6&"_"&$G$7&"_{q}"')
        put(ws, f"{col}41", f"=$G$8*{cg}", "0.0%", CALC)
        put(ws, f"{col}42", f"=$B$25*(1+{col}41)", "#,##0", CALC)
        put(ws, f"{col}50", f"=$B$11*{col}41", "#,##0", CALC)
    put(ws, "B42", "=$B$25", "#,##0", CALC)
    for col in "BCDE":
        put(ws, f"{col}43", f"=ROUNDUP({col}42/$B$26,0)", "#,##0", CALC)
        put(ws, f"{col}44", f"={col}43/$B$27", "0.0", CALC)
        put(ws, f"{col}45", f"={col}43-$B$43", "#,##0", CALC)
        put(ws, f"{col}46", f"={col}43*$B$28", "#,##0", CALC)
        put(ws, f"{col}47", f"=ROUNDUP({col}46/$B$30,0)", "#,##0", CALC)
        put(ws, f"{col}48", f"={col}46*$B$29", "#,##0", CALC)
        put(ws, f"{col}49", f"={col}42*$B$31", "$#,##0", CALC)
    put(ws, "A52", "How to choose between 80% and 95%", bold=True)
    put(ws, "A53", '="At 80%, you accept that a session of patients turned away costs about "&ROUND(0.8/0.2,0)&" times what an idle session costs. At 95%, about "&ROUND(0.95/0.05,0)&" times."')
    put(ws, "A54", '="Going from 80% to 95% adds "&(E43-D43)&" sessions and "&TEXT(E48-D48,"#,##0")&" volunteer hours a year."')
    put(ws, "A56", "Check", bold=True)
    put(ws, "A57", '="Your patients are "&TEXT(B11/B6,"0.0%")&" of the county\'s uninsured low-income adults. "&IF(B11/B6>0.25,"That is high; check that the county and patient counts match.",IF(B11/B6<0.01,"That is low; the planner assumes you draw new patients in proportion to your current share.","That is a plausible share."))')
    for r_ in (53, 54, 57):
        ws.cell(r_, 1).alignment = WRAP
        ws.merge_cells(start_row=r_, start_column=1, end_row=r_, end_column=5)
        ws.row_dimensions[r_].height = 32
    dv_s = DataValidation(type="list", formula1='"Lower demand,Central,Higher demand"', allow_blank=False)
    dv_y = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    dv_c = DataValidation(type="list", formula1=f"={label_rng}", allow_blank=False)
    dv_r = DataValidation(type="decimal", operator="between", formula1="0", formula2="5")
    for dv, ref in ((dv_s, "B34"), (dv_y, "B35"), (dv_c, "B5"), (dv_r, "B36")):
        ws.add_data_validation(dv)
        dv.add(ref)
    ws.freeze_panes = "A4"

    # ---------------- Need statement
    ws = need
    ws.column_dimensions["A"].width = 110
    ws["A1"] = "Need statement for a state or foundation application"
    ws["A1"].font = TITLE
    ws["A2"] = "Built from the Planner sheet. Copy the paragraphs into your application and edit freely."
    p = "Planner!"
    key = f'"add_"&{p}$G$6&"_"&{p}$G$7&"_'
    def add(q):
        return look(f'{key}{q}"', f"{p}$G$5")
    paras = [
        f'="Need. In "&{p}B5&", about "&TEXT({p}B6,"#,##0")&" adults with incomes at or below 200% of the federal poverty level were uninsured in 2024 (Census Small Area Health Insurance Estimates). '
        f'We estimate that between "&TEXT({add("p10")},"#,##0")&" and "&TEXT({add("p90")},"#,##0")&" more adults (an 80% interval: the 10th to 90th percentile of our scenario draws) '
        f'will be uninsured in 2027 than in 2024, counting losses of paid marketplace coverage from February 2025 on (central estimate "&TEXT({add("p50")},"#,##0")&"), '
        f'a rise of "&TEXT({p}C40,"0%")&" at the central estimate."',
        f'="Our capacity. In the last 12 months our clinic provided "&TEXT({p}B25,"#,##0")&" visits to "&TEXT({p}B11,"#,##0")&" uninsured adults, in "&TEXT({p}B43,"#,##0")&" sessions. '
        f'To meet 2027 demand with 80% probability we plan "&TEXT({p}D43,"#,##0")&" sessions ("&TEXT({p}D45,"#,##0")&" more than today), "&TEXT({p}D47,"#,##0")&" volunteers and about "&TEXT({p}D49,"$#,##0")&" in medication. '
        f'With 95% probability, "&TEXT({p}E43,"#,##0")&" sessions, "&TEXT({p}E47,"#,##0")&" volunteers and "&TEXT({p}E49,"$#,##0")&"."',
        # No references, so plain text: Excel rejects a text literal over 255 characters inside a formula.
        "Method and limits. The range comes from public data (CMS Open Enrollment files for 2025 and 2026, CMS effectuated enrollment, Census SAHIE 2024, Medicaid.gov enrollment) "
        "and stated assumptions about how many people lose coverage and end up uninsured. It is a planning range, not a forecast, and it has not been checked against 2027 outcomes. "
        "The interval and the session counts hold within the scenario chosen on the Planner sheet; they do not include the spread between scenarios. "
        "Plan selections are not paid coverage. We will replace the assumption about new arrivals with our own intake counts.",
    ]
    for i, f in enumerate(paras):
        c = ws.cell(4 + 2 * i, 1, f)
        c.alignment = WRAP
        ws.row_dimensions[4 + 2 * i].height = 82

    # ---------------- Assumptions
    ws = assume
    ws.column_dimensions["A"].width = 52
    for col in "BCDE":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["F"].width = 70
    ws["A1"] = "What is data and what is assumption"
    ws["A1"].font = TITLE
    rows = [
        ("Data (observed, public)", None, None, None, None, None),
        ("Low-income plan selections, 2025 and 2026 (CMS OEP county files)", "", "", "", "Administrative counts of everyone who selected a plan, not a sample. Suppressed cells (1 to 10 people) are redrawn at random.", None),
        ("Share of selections paid up in February (CMS effectuated enrollment)", "", "", "", "State level, HealthCare.gov platform, applied to each county in the state.", None),
        ("Uninsured adults at or below 200% of poverty, 2024 (Census SAHIE)", "", "", "", "Modeled estimate with 90% margins of error, which the simulation draws from.", None),
        ("Medicaid and CHIP enrollment by state, June 2026 (Medicaid.gov)", "", "", "", "Used only for the Medicaid layer, split to counties by adults at or below 138% of poverty.", None),
        ("Assumptions (scenario means)", "Lower", "Central", "Higher", "Spread", "What it means"),
    ]
    names = [("u", "Lost marketplace coverage that ends in no coverage", "Beta, concentration {u}"),
             ("d", "Further loss of paid marketplace coverage after Feb 2026", "Beta, concentration {d}"),
             ("w", "Medicaid expansion adults who eventually lose coverage", "Beta, concentration {w}; from our Medicaid what-if post"),
             ("z", "Those who lose Medicaid and end up uninsured", "Beta, concentration {z}"),
             ("r", "Share of the Medicaid loss in place on average in 2027", "Beta, concentration {r}")]
    for r_, row in enumerate(rows, start=3):
        for c_, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r_, c_, v)
        if row[1] is None or row[1] == "Lower":
            for c_ in range(1, 7):
                ws.cell(r_, c_).fill = HEAD
                ws.cell(r_, c_).font = BOLD
    r_ = 3 + len(rows)
    for k, label, spread in names:
        ws.cell(r_, 1, label)
        for c_, s in enumerate(("lower", "central", "higher"), start=2):
            ws.cell(r_, c_, model.SCENARIOS[s][k]).number_format = "0.0%"
        ws.cell(r_, 5, spread.format(**model.KAPPA))
        r_ += 1
    extra = [
        ("Expansion adults as a share of Medicaid/CHIP enrollment", model.E_MEAN, "KFF reports 31% of Medicaid enrollment in expansion states (June 2024). We apply it to Medicaid plus CHIP enrollment, a larger base, so use 30%. Beta, concentration 60."),
        ("Share of the 2026 marketplace loss already in your trailing-year visits", model.PHI, "A timing assumption, fixed in every scenario. It trims the clinic's growth (row 41 of the Planner) but not the county pool, which counts the whole loss."),
        ("Newly uninsured adults come to you at your current rate per uninsured adult", 1.0, "The weakest link and the cell you overwrite (B36 on the Planner). Replace it with your own intake counts."),
    ]
    for label, v, note in extra:
        ws.cell(r_, 1, label)
        ws.cell(r_, 3, v).number_format = "0.0%" if v < 1 else "0.00"
        ws.cell(r_, 6, note)
        r_ += 1
    ws.cell(r_ + 1, 1, "Each scenario runs 10,000 draws. Parameters are shared across counties in a draw, so uncertainty does not average out between counties.").alignment = WRAP
    ws.cell(r_ + 2, 1, "Selections are not paid coverage, and the file covers 30 HealthCare.gov states only (not California, New York, Illinois and other state-run exchanges).").alignment = WRAP
    ws.cell(r_ + 3, 1, "Nothing here is a causal effect of any policy. The demand range cannot be checked against outcomes with these data.").alignment = WRAP
    for rr in range(3, r_ + 4):
        ws.cell(rr, 6).alignment = WRAP
        ws.cell(rr, 5).alignment = WRAP

    # ---------------- Intake guide
    ws = intake
    ws.column_dimensions["A"].width = 120
    ws["A1"] = "Three numbers to log at intake"
    ws["A1"].font = TITLE
    lines = [
        "1. Coverage at the first visit, and for people who lost coverage, the month it ended. Count new uninsured patients each month. This replaces the weakest assumption in the planner.",
        "2. Household income as a share of the poverty level: at or below 138%, 139% to 200%, above 200%. The planner covers 200% and below.",
        "3. New or returning, and the number of visits in the first six months for new patients. It tells you whether new arrivals need more or fewer sessions than current patients.",
        "",
        "How to use them. After three months of 2027, compare your count of new uninsured patients with the 'New patients expected' row on the Planner (scaled to three months).",
        "If you are running above the 80% column, switch to the Higher demand scenario and rerun. If you are below the Likely column, switch to Lower demand.",
        "Counts only. Do not record names, dates of birth or any other identifier in this workbook.",
    ]
    for i, t in enumerate(lines, start=3):
        ws.cell(i, 1, t).alignment = WRAP
        ws.row_dimensions[i].height = 34

    # ---------------- Start here
    ws = start
    ws.column_dimensions["A"].width = 110
    ws["A1"] = "Free Clinic Demand Planner"
    ws["A1"].font = TITLE
    intro = [
        "An example from Primordial Analytics on public data. It is not a client engagement and the clinic in it is hypothetical.",
        "",
        "What it does. It gives a planning range for how many more low-income uninsured adults your county may have in 2027, and turns it into sessions, volunteers and medication dollars at 80% or 95% probability.",
        "",
        "How to use it.",
        "1. Open the Planner sheet. Pick your county in cell B5. Only the 30 HealthCare.gov states are here, and Illinois is not.",
        "2. Replace the yellow cells in section 2 with your own numbers: 12 months of visits, patients per session, volunteers per session.",
        "3. Choose a scenario and decide whether to include the Medicaid work-requirement layer (a labelled placeholder).",
        "4. Read section 4. Copy the Need statement sheet into your application.",
        "",
        "What it cannot do. It is not a forecast. Plan selections are not paid coverage, the share of new arrivals who come to a free clinic is an assumption, and the Medicaid layer rests on placeholder loss rates. See the Assumptions sheet.",
        "No macros, no external links. Everything is a formula over the Counties sheet. Free to use, copy and share.",
        "Open it in Excel, LibreOffice or Numbers. A previewer that does not calculate (some mail and phone viewers) shows blank result cells.",
        "",
        "Code and method: https://github.com/primordial-analytics/primordial-analytics.github.io/tree/master/analysis/11-free-clinic-surge-planner",
    ]
    for i, t in enumerate(intro, start=3):
        ws.cell(i, 1, t).alignment = WRAP
        if t:
            ws.row_dimensions[i].height = 32
    wb.calculation.fullCalcOnLoad = True
    out_dir.mkdir(exist_ok=True)
    path = out_dir / "11-free-clinic-surge-planner-demand-planner.xlsx"
    wb.save(path)
    add_cached_values(path)
